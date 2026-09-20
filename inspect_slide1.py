import pptx

prs = pptx.Presentation(r'C:\Users\yashc\Downloads\ppp.pptx')
slide1 = prs.slides[0]
for idx, shape in enumerate(slide1.shapes):
    if shape.has_text_frame:
        print(f"Shape {idx} (Name: {shape.name}):")
        for p_idx, p in enumerate(shape.text_frame.paragraphs):
            print(f"  P{p_idx}: '{p.text}'")
            for r_idx, r in enumerate(p.runs):
                print(f"    Run {r_idx}: '{r.text}'")
