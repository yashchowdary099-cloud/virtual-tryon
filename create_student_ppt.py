import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_student_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5) # 4:3 ratio matching standard college templates

    # Standard template color palette
    NAVY = RGBColor(16, 44, 87)       # #102C57 Dark Navy
    RED = RGBColor(180, 30, 30)       # Logo Header Red
    TEXT_DARK = RGBColor(30, 30, 30)
    BG_WHITE = RGBColor(255, 255, 255)
    GRAY_BORDER = RGBColor(200, 200, 200)
    LIGHT_BOX = RGBColor(240, 244, 248)
    HEADER_BLUE = RGBColor(24, 76, 120)

    # Module box accent colors (like in friend's slide)
    BLUE_ACCENT = RGBColor(41, 128, 185)
    GREEN_ACCENT = RGBColor(39, 174, 96)
    ORANGE_ACCENT = RGBColor(230, 126, 34)
    PURPLE_ACCENT = RGBColor(142, 68, 173)
    TEAL_ACCENT = RGBColor(22, 160, 133)
    RED_ACCENT = RGBColor(192, 57, 43)
    DARK_BLUE_ACCENT = RGBColor(44, 62, 80)

    blank_layout = prs.slide_layouts[6]

    def apply_slide_frame(slide, title_text="", slide_num=1, is_title_slide=False):
        # White background
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = BG_WHITE

        # Outer Dark Blue Frame Line around slide
        outer_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.3), Inches(0.3), Inches(9.4), Inches(6.9))
        outer_box.fill.background()
        outer_box.line.color.rgb = NAVY
        outer_box.line.width = Pt(1.5)

        if is_title_slide:
            # Sathyabama Header Logo Text
            h_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9.0), Inches(1.5))
            tf = h_box.text_frame
            tf.word_wrap = True
            tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
            
            p1 = tf.paragraphs[0]
            p1.text = "SATHYABAMA"
            p1.font.bold = True
            p1.font.size = Pt(28)
            p1.font.color.rgb = RED
            p1.alignment = PP_ALIGN.CENTER

            p2 = tf.add_paragraph()
            p2.text = "INSTITUTE OF SCIENCE AND TECHNOLOGY\n(DEEMED TO BE UNIVERSITY)"
            p2.font.bold = True
            p2.font.size = Pt(12)
            p2.font.color.rgb = RED
            p2.alignment = PP_ALIGN.CENTER

            p3 = tf.add_paragraph()
            p3.text = "CATEGORY - 1 UNIVERSITY BY UGC\nAccredited with 'A++' grade by NAAC | Approved by AICTE\nwww.sathyabama.ac.in"
            p3.font.size = Pt(9)
            p3.font.color.rgb = NAVY
            p3.alignment = PP_ALIGN.CENTER

            # Divider Red Line under header
            div_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(2.0), Inches(9.2), Inches(0.02))
            div_line.fill.solid()
            div_line.fill.fore_color.rgb = RED
            div_line.line.fill.background()
        elif title_text:
            # Slide Header Title
            title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.45), Inches(9.0), Inches(0.8))
            tf = title_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title_text
            p.font.bold = True
            p.font.size = Pt(24)
            p.font.color.rgb = NAVY
            p.alignment = PP_ALIGN.CENTER

            # Line under slide title
            div_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(1.25), Inches(9.2), Inches(0.02))
            div_line.fill.solid()
            div_line.fill.fore_color.rgb = NAVY
            div_line.line.fill.background()

        # Footer Line & Text
        f_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(6.8), Inches(9.2), Inches(0.01))
        f_line.fill.solid()
        f_line.fill.fore_color.rgb = GRAY_BORDER
        f_line.line.fill.background()

        # Left Footer (Date)
        date_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.85), Inches(2.5), Inches(0.35))
        p_d = date_box.text_frame.paragraphs[0]
        p_d.text = "20 September 2026"
        p_d.font.size = Pt(9)
        p_d.font.color.rgb = GRAY_BORDER

        # Center Footer (School Name)
        school_box = slide.shapes.add_textbox(Inches(3.2), Inches(6.85), Inches(3.6), Inches(0.35))
        p_s = school_box.text_frame.paragraphs[0]
        p_s.text = "School of Computing - CSE"
        p_s.font.size = Pt(9)
        p_s.font.color.rgb = GRAY_BORDER
        p_s.alignment = PP_ALIGN.CENTER

        # Right Footer (Slide Number)
        num_box = slide.shapes.add_textbox(Inches(8.3), Inches(6.85), Inches(1.2), Inches(0.35))
        p_n = num_box.text_frame.paragraphs[0]
        p_n.text = str(slide_num)
        p_n.font.size = Pt(9)
        p_n.font.color.rgb = GRAY_BORDER
        p_n.alignment = PP_ALIGN.RIGHT

    # =============================================================
    # Slide 1: Title Slide (Updated Student & Guide Details)
    # =============================================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_slide_frame(s1, slide_num=1, is_title_slide=True)

    t_box1 = s1.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(9.0), Inches(2.5))
    tf1 = t_box1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING"
    p.font.bold = True
    p.font.size = Pt(17)
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER

    p = tf1.add_paragraph()
    p.text = "Professional Training I"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_DARK
    p.alignment = PP_ALIGN.CENTER

    p = tf1.add_paragraph()
    p.text = "\nEmployee Attendance & HR Management System"
    p.font.bold = True
    p.font.size = Pt(24)
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER

    # Student & Guide Details Block
    st_box = s1.shapes.add_textbox(Inches(0.6), Inches(5.0), Inches(4.2), Inches(1.6))
    tf_st = st_box.text_frame
    tf_st.word_wrap = True

    p = tf_st.paragraphs[0]
    p.text = "PROJECT STUDENT"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = NAVY

    p = tf_st.add_paragraph()
    p.text = "NAGABHIRU YASHWANTH SAI, 44111243"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK

    p = tf_st.add_paragraph()
    p.text = "CSE"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK

    gd_box = s1.shapes.add_textbox(Inches(5.0), Inches(5.0), Inches(4.4), Inches(1.6))
    tf_gd = gd_box.text_frame
    tf_gd.word_wrap = True

    p = tf_gd.paragraphs[0]
    p.text = "GUIDE"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = NAVY

    p = tf_gd.add_paragraph()
    p.text = "Ms. R. Velvizhi,"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK

    p = tf_gd.add_paragraph()
    p.text = "Assistant Professor, CSE"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK

    # =============================================================
    # Slide 2: AGENDA
    # =============================================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_slide_frame(s2, title_text="AGENDA", slide_num=2)

    ag_box = s2.shapes.add_textbox(Inches(1.5), Inches(1.6), Inches(7.0), Inches(5.0))
    tf2 = ag_box.text_frame
    tf2.word_wrap = True

    agenda_items = [
        "Certificate",
        "Abstract",
        "Introduction",
        "Objective",
        "System Architecture / Ideation Map",
        "Module Implementation",
        "Results and Discussions",
        "Conclusion",
        "References"
    ]

    for i, item in enumerate(agenda_items):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = f"•   {item}"
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(10)

    # =============================================================
    # Slide 3: COURSE CERTIFICATE
    # =============================================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_slide_frame(s3, title_text="COURSE CERTIFICATE", slide_num=3)

    # Blank placeholder frame matching template page 3
    cert_frame = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(1.8), Inches(7.0), Inches(4.5))
    cert_frame.fill.solid()
    cert_frame.fill.fore_color.rgb = BG_WHITE
    cert_frame.line.color.rgb = GRAY_BORDER

    # =============================================================
    # Slide 4: Abstract (PARAGRAPH FORMAT AS REQUESTED - NO BULLETS)
    # =============================================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_slide_frame(s4, title_text="Abstract", slide_num=4)

    ab_box = s4.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(8.8), Inches(5.0))
    tf4 = ab_box.text_frame
    tf4.word_wrap = True

    abstract_text = (
        "The Employee Attendance & HR Management System is a cloud-based application "
        "designed to simplify employee and human resource activities. It provides features "
        "such as employee registration, secure login, attendance tracking, leave management, "
        "employee records, and report generation. The application is deployed on AWS EC2 "
        "for reliable access and operation. Amazon RDS stores employee and attendance data, "
        "while Amazon S3 securely stores documents and files. IAM and VPC provide access "
        "control and network security. Elastic Load Balancer and Auto Scaling improve availability "
        "and scalability. CloudWatch monitors system performance and resource usage. The proposed "
        "system provides a secure, reliable, scalable, and efficient solution for modern employee "
        "attendance and HR management."
    )

    p = tf4.paragraphs[0]
    p.text = abstract_text
    p.font.size = Pt(15)
    p.font.color.rgb = TEXT_DARK
    p.line_spacing = 1.35

    # =============================================================
    # Slide 5: Introduction
    # =============================================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_slide_frame(s5, title_text="Introduction", slide_num=5)

    in_box = s5.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(8.8), Inches(5.0))
    tf5 = in_box.text_frame
    tf5.word_wrap = True

    intro_bullets = [
        "Employee attendance and HR management are essential for maintaining accurate employee records and organizational efficiency.",
        "Manual attendance and HR processes are time-consuming and can lead to errors.",
        "The proposed system provides digital features for attendance, employee records, leave management, and report generation.",
        "The application is deployed on AWS EC2, with RDS, S3, IAM, VPC, ELB, Auto Scaling, and CloudWatch supporting storage, security, scalability, and monitoring."
    ]

    for i, b in enumerate(intro_bullets):
        p = tf5.paragraphs[0] if i == 0 else tf5.add_paragraph()
        p.text = f"•  {b}"
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(16)
        p.line_spacing = 1.25

    # =============================================================
    # Slide 6: Objective
    # =============================================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_slide_frame(s6, title_text="Objective", slide_num=6)

    ob_box = s6.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(8.8), Inches(5.0))
    tf6 = ob_box.text_frame
    tf6.word_wrap = True

    objectives = [
        "To develop a centralized Employee Attendance and HR Management System.",
        "To automate employee attendance, leave, and record management.",
        "To securely store employee data using cloud database services.",
        "To deploy the application on AWS EC2 with scalability and high availability.",
        "To monitor and manage AWS resources using IAM, VPC, ELB, Auto Scaling, and CloudWatch."
    ]

    for i, obj in enumerate(objectives):
        p = tf6.paragraphs[0] if i == 0 else tf6.add_paragraph()
        p.text = f"•  {obj}"
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(14)
        p.line_spacing = 1.25

    # =============================================================
    # Slide 7: System Architecture / Ideation Map (Human Student Style Diagram)
    # =============================================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_slide_frame(s7, title_text="System Architecture / Ideation Map", slide_num=7)

    # Top Title Banner inside slide
    banner = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.2), Inches(1.4), Inches(5.6), Inches(0.5))
    banner.fill.solid()
    banner.fill.fore_color.rgb = NAVY
    banner.line.fill.background()
    p = banner.text_frame.paragraphs[0]
    p.text = "Employee Attendance & HR Management System\nSystem Architecture / Ideation Map"
    p.font.bold = True
    p.font.size = Pt(10)
    p.font.color.rgb = BG_WHITE
    p.alignment = PP_ALIGN.CENTER

    # Left Side: Users Box
    users_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2.1), Inches(1.6), Inches(3.2))
    users_box.fill.solid()
    users_box.fill.fore_color.rgb = LIGHT_BOX
    users_box.line.color.rgb = BLUE_ACCENT
    users_box.line.width = Pt(1.5)

    tf_u = users_box.text_frame
    tf_u.word_wrap = True
    p = tf_u.paragraphs[0]
    p.text = "USERS\n"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER

    user_types = ["• Admin\n  (Manage System)", "• HR\n  (Manage Employees)", "• Employee\n  (Attendance / Leave)", "• Others\n  (View Reports)"]
    for ut in user_types:
        p = tf_u.add_paragraph()
        p.text = ut
        p.font.size = Pt(9)
        p.font.color.rgb = TEXT_DARK

    # Internet Circle Connector
    net_circle = s7.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2.25), Inches(3.3), Inches(0.8), Inches(0.8))
    net_circle.fill.solid()
    net_circle.fill.fore_color.rgb = BLUE_ACCENT
    net_circle.line.fill.background()
    p = net_circle.text_frame.paragraphs[0]
    p.text = "Internet"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = BG_WHITE
    p.alignment = PP_ALIGN.CENTER

    # Middle Main: AWS Cloud Box
    cloud_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.2), Inches(2.1), Inches(4.6), Inches(3.2))
    cloud_box.fill.solid()
    cloud_box.fill.fore_color.rgb = BG_WHITE
    cloud_box.line.color.rgb = NAVY
    cloud_box.line.width = Pt(1.5)

    p = cloud_box.text_frame.paragraphs[0]
    p.text = "AWS Cloud (VPC Isolated Network)"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.LEFT

    # Inside Cloud: ELB Box
    elb = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(2.5), Inches(2.0), Inches(0.4))
    elb.fill.solid()
    elb.fill.fore_color.rgb = PURPLE_ACCENT
    elb.line.fill.background()
    p = elb.text_frame.paragraphs[0]
    p.text = "Elastic Load Balancer"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = BG_WHITE
    p.alignment = PP_ALIGN.CENTER

    # EC2 Zone 1 & Zone 2
    ec2_1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.5), Inches(3.2), Inches(1.4), Inches(0.7))
    ec2_1.fill.solid()
    ec2_1.fill.fore_color.rgb = ORANGE_ACCENT
    ec2_1.line.fill.background()
    p = ec2_1.text_frame.paragraphs[0]
    p.text = "EC2 Web Server\n(Zone 1)"
    p.font.size = Pt(9)
    p.font.color.rgb = BG_WHITE
    p.alignment = PP_ALIGN.CENTER

    ec2_2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.2), Inches(3.2), Inches(1.4), Inches(0.7))
    ec2_2.fill.solid()
    ec2_2.fill.fore_color.rgb = ORANGE_ACCENT
    ec2_2.line.fill.background()
    p = ec2_2.text_frame.paragraphs[0]
    p.text = "EC2 Web Server\n(Zone 2)"
    p.font.size = Pt(9)
    p.font.color.rgb = BG_WHITE
    p.alignment = PP_ALIGN.CENTER

    # Auto Scaling in Middle
    asg = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.0), Inches(3.3), Inches(1.1), Inches(0.5))
    asg.fill.solid()
    asg.fill.fore_color.rgb = LIGHT_BOX
    asg.line.color.rgb = ORANGE_ACCENT
    p = asg.text_frame.paragraphs[0]
    p.text = "Auto Scaling"
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p.alignment = PP_ALIGN.CENTER

    # DB & S3 Box
    rds = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.6), Inches(4.3), Inches(1.7), Inches(0.7))
    rds.fill.solid()
    rds.fill.fore_color.rgb = GREEN_ACCENT
    rds.line.fill.background()
    p = rds.text_frame.paragraphs[0]
    p.text = "Amazon RDS\n(Employee Database)"
    p.font.size = Pt(8)
    p.font.color.rgb = BG_WHITE
    p.alignment = PP_ALIGN.CENTER

    s3_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.9), Inches(4.3), Inches(1.7), Inches(0.7))
    s3_box.fill.solid()
    s3_box.fill.fore_color.rgb = GREEN_ACCENT
    s3_box.line.fill.background()
    p = s3_box.text_frame.paragraphs[0]
    p.text = "Amazon S3\n(Document Storage)"
    p.font.size = Pt(8)
    p.font.color.rgb = BG_WHITE
    p.alignment = PP_ALIGN.CENTER

    # Right Side: IAM & CloudWatch
    iam = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(2.3), Inches(1.5), Inches(1.3))
    iam.fill.solid()
    iam.fill.fore_color.rgb = LIGHT_BOX
    iam.line.color.rgb = RED_ACCENT
    p = iam.text_frame.paragraphs[0]
    p.text = "IAM\nAccess Control &\nPermissions"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = RED_ACCENT
    p.alignment = PP_ALIGN.CENTER

    cw = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(3.9), Inches(1.5), Inches(1.3))
    cw.fill.solid()
    cw.fill.fore_color.rgb = LIGHT_BOX
    cw.line.color.rgb = PURPLE_ACCENT
    p = cw.text_frame.paragraphs[0]
    p.text = "CloudWatch\nMonitoring &\nSystem Logs"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = PURPLE_ACCENT
    p.alignment = PP_ALIGN.CENTER

    # Bottom Row: Key Features Bar
    ft_bar = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(5.5), Inches(9.0), Inches(0.9))
    ft_bar.fill.solid()
    ft_bar.fill.fore_color.rgb = LIGHT_BOX
    ft_bar.line.color.rgb = BLUE_ACCENT
    p = ft_bar.text_frame.paragraphs[0]
    p.text = "System Modules: Employee Registration | Secure Login | Attendance Tracking | Leave Management | Employee Records | Reports | Document Storage"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER

    # =============================================================
    # Slide 8: Module Implementation (7 Clean Colored Blocks)
    # =============================================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_slide_frame(s8, title_text="Module Implementation", slide_num=8)

    modules = [
        ("1. User Management", BLUE_ACCENT, ["User Registration", "Login / Logout", "Role-based Access", "Profile Management"]),
        ("2. Attendance Tracking", GREEN_ACCENT, ["Check-in / Check-out", "View History", "Leave Balance Update", "Mark via Web/Mobile"]),
        ("3. Leave Management", ORANGE_ACCENT, ["Apply for Leave", "Approve / Reject", "Leave Balance", "Leave History"]),
        ("4. Employee Records", PURPLE_ACCENT, ["Add/Update Employees", "Department Mgmt", "Designation Mgmt", "Employee Details View"]),
        ("5. Reports", TEAL_ACCENT, ["Attendance Report", "Leave Report", "Employee Report", "Export (PDF/Excel)"]),
        ("6. Document Storage", RED_ACCENT, ["Store Documents", "Access Files", "Secure Storage", "Manage Permissions"]),
        ("7. Monitoring & Security", DARK_BLUE_ACCENT, ["IAM Access Control", "VPC Network Security", "Load Balancer Monitor", "CloudWatch Logs"])
    ]

    box_w = Inches(1.28)
    box_h = Inches(4.8)

    for idx, (m_title, m_color, m_items) in enumerate(modules):
        left_pos = Inches(0.5 + idx * 1.32)
        top_pos = Inches(1.5)

        m_box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, box_w, box_h)
        m_box.fill.solid()
        m_box.fill.fore_color.rgb = LIGHT_BOX
        m_box.line.color.rgb = m_color
        m_box.line.width = Pt(1.5)

        tf = m_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.04)

        # Header Box inside
        hdr = m_box.text_frame.paragraphs[0]
        hdr.text = m_title
        hdr.font.bold = True
        hdr.font.size = Pt(9)
        hdr.font.color.rgb = m_color
        hdr.alignment = PP_ALIGN.CENTER
        hdr.space_after = Pt(8)

        for item in m_items:
            p = tf.add_paragraph()
            p.text = f"• {item}"
            p.font.size = Pt(8)
            p.font.color.rgb = TEXT_DARK
            p.space_after = Pt(4)

    # =============================================================
    # Slide 9: Results and Discussions
    # =============================================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_slide_frame(s9, title_text="Results and Discussions", slide_num=9)

    res_box = s9.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(8.8), Inches(5.0))
    tf9 = res_box.text_frame
    tf9.word_wrap = True

    results = [
        "The system was successfully developed and deployed on AWS EC2 for easy access.",
        "Amazon RDS securely stores employee and attendance information.",
        "Amazon S3 provides reliable storage for documents and files.",
        "IAM, VPC, ELB, and Auto Scaling improve security, availability, and scalability.",
        "CloudWatch monitors performance, while the overall system reduces manual work and improves HR management efficiency."
    ]

    for i, r in enumerate(results):
        p = tf9.paragraphs[0] if i == 0 else tf9.add_paragraph()
        p.text = f"•  {r}"
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(14)
        p.line_spacing = 1.25

    # =============================================================
    # Slide 10: Conclusion
    # =============================================================
    s10 = prs.slides.add_slide(blank_layout)
    apply_slide_frame(s10, title_text="Conclusion", slide_num=10)

    conc_box = s10.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(8.8), Inches(5.0))
    tf10 = conc_box.text_frame
    tf10.word_wrap = True

    conclusions = [
        "The Employee Attendance & HR Management System provides an efficient solution for managing employee activities.",
        "AWS services enable secure data storage, reliable application deployment, and easy accessibility.",
        "The system reduces manual work and improves attendance and HR management.",
        "Security, scalability, and performance are improved using AWS services such as IAM, VPC, ELB, Auto Scaling, and CloudWatch.",
        "Overall, the project provides a secure, reliable, scalable, and cloud-based HR management solution."
    ]

    for i, c in enumerate(conclusions):
        p = tf10.paragraphs[0] if i == 0 else tf10.add_paragraph()
        p.text = f"•  {c}"
        p.font.size = Pt(15)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(14)
        p.line_spacing = 1.25

    # =============================================================
    # Slide 11: REFERENCES
    # =============================================================
    s11 = prs.slides.add_slide(blank_layout)
    apply_slide_frame(s11, title_text="REFERENCES", slide_num=11)

    ref_box = s11.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(8.8), Inches(5.2))
    tf11 = ref_box.text_frame
    tf11.word_wrap = True

    references = [
        "• 1]. Kavis, M. J. Architecting the Cloud: Design Decisions for Cloud Computing Service Models (SaaS, PaaS, and IaaS). Wiley (2014).",
        "• 2]. Erl, T., Puttini, R. & Mahmood, Z. Cloud Computing: Concepts, Technology & Architecture. Pearson (2013).",
        "• 3]. Sosinsky, B. Cloud Computing Bible. Wiley (2011).",
        "• 4]. Vacca, J. R. Cloud Computing Security: Foundations and Challenges. CRC Press (2017).",
        "• 5]. Kurose, J. F. & Ross, K. W. Computer Networking: A Top-Down Approach. Pearson (2021).",
        "• 6]. Elmasri, R. & Navathe, S. B. Fundamentals of Database Systems. Pearson (2016).",
        "• 7]. Sommerville, I. Software Engineering. Pearson (2015).",
        "• 8]. Bass, L., Clements, P. & Kazman, R. Software Architecture in Practice. Addison-Wesley Professional (2021)."
    ]

    for i, ref in enumerate(references):
        p = tf11.paragraphs[0] if i == 0 else tf11.add_paragraph()
        p.text = ref
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # =============================================================
    # Slide 12: THANK YOU
    # =============================================================
    s12 = prs.slides.add_slide(blank_layout)
    apply_slide_frame(s12, title_text="THANK YOU", slide_num=12)

    ty_box = s12.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(8.4), Inches(3.5))
    tf12 = ty_box.text_frame
    tf12.word_wrap = True

    p = tf12.paragraphs[0]
    p.text = "We thank God, Our Department, Guide, Panel Members, Supportive Professors and all Technical and non Technical staff who helped us in our Project."
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(20)

    # Save output to desktop
    output_path = r"c:\Users\yashc\OneDrive\Desktop\Employee_Attendance_HR_Management_System.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    build_student_presentation()
